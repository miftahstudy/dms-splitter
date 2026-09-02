import io
import os
import re
import math
import zipfile
from datetime import datetime, date, timedelta

import pandas as pd
import streamlit as st
import openpyxl

# ---------------------------------------------------------------------------
# Konfigurasi kolom sesuai template DMS
# ---------------------------------------------------------------------------
COLUMNS = [
    {"key": "customer_code", "label": "Customer Code", "mandatory": True, "tokens": ["customer code"]},
    {"key": "customer_name", "label": "Customer Name", "mandatory": False, "tokens": ["customer name"]},
    {"key": "branch_code", "label": "Customer Branch Code", "mandatory": True, "tokens": ["branch code"]},
    {"key": "branch_name", "label": "Customer Branch Name", "mandatory": False, "tokens": ["branch name"]},
    {"key": "address", "label": "Customer Address", "mandatory": False, "tokens": ["customer address", "address"]},
    {"key": "po_date", "label": "PO Date", "mandatory": True, "tokens": ["po date"]},
    {"key": "po_number", "label": "PO Number", "mandatory": True, "tokens": ["po number"]},
    {"key": "store_code", "label": "Customer Store Code", "mandatory": True, "tokens": ["store code"]},
    {"key": "store_name", "label": "Customer Store Name", "mandatory": False, "tokens": ["store name"]},
    {"key": "sku_code", "label": "Customer SKU Code", "mandatory": True, "tokens": ["sku code"]},
    {"key": "sku_name", "label": "Customer SKU Name", "mandatory": False, "tokens": ["sku name"]},
    {"key": "qty", "label": "Qty", "mandatory": True, "tokens": ["qty", "quantity"]},
]
COL_KEYS = [c["key"] for c in COLUMNS]
MANDATORY_TEXT_KEYS = [c["key"] for c in COLUMNS if c["mandatory"] and c["key"] not in ("po_date", "qty")]
LABEL_BY_KEY = {c["key"]: c["label"] for c in COLUMNS}

DEFAULT_TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "assets", "default_template.xlsx")

BULAN_ID = {
    "januari": "january", "februari": "february", "maret": "march", "april": "april",
    "mei": "may", "juni": "june", "juli": "july", "agustus": "august",
    "september": "september", "oktober": "october", "november": "november", "desember": "december",
    "jan": "jan", "feb": "feb", "mar": "mar", "apr": "apr", "jun": "jun", "jul": "jul",
    "agu": "aug", "ags": "aug", "okt": "oct", "nov": "nov", "des": "dec",
}

EXPLICIT_DATE_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d",
    "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%Y.%m.%d",
    "%d-%m-%y", "%d/%m/%y",
    "%m/%d/%Y", "%m-%d-%Y",
    "%d %B %Y", "%d %b %Y", "%B %d, %Y", "%b %d, %Y",
    "%d-%b-%Y", "%d-%B-%Y",
]

VALID_DATE_STRING_RE = re.compile(r"^\d{4}[-/]\d{2}[-/]\d{2}$")


def norm(s):
    if s is None:
        return ""
    s = str(s).replace("\n", " ").replace("\r", " ").replace("*", "")
    s = " ".join(s.split())
    return s.strip().lower()


def find_header_row(rows, max_scan=10):
    best_idx, best_score = 0, -1
    for i, row in enumerate(rows[:max_scan]):
        normed = [norm(v) for v in row]
        score = 0
        for col in COLUMNS:
            for cell in normed:
                if any(tok in cell for tok in col["tokens"]):
                    score += 1
                    break
        if score > best_score:
            best_score = score
            best_idx = i
    return best_idx, best_score


def auto_map_columns(header_row_values):
    normed = [norm(v) for v in header_row_values]
    mapping = {}
    used = set()
    for col in COLUMNS:
        found = None
        for idx, cell in enumerate(normed):
            if idx in used:
                continue
            if any(tok in cell for tok in col["tokens"]):
                found = idx
                break
        if found is not None:
            mapping[col["key"]] = found
            used.add(found)
    return mapping


def excel_serial_to_date(value):
    try:
        d = datetime(1899, 12, 30) + timedelta(days=float(value))
        if 1990 <= d.year <= 2100:
            return d.date()
    except Exception:
        pass
    return None


def try_parse_date(value):
    if value is None:
        return None, "kosong"
    if isinstance(value, datetime):
        return value.date(), None
    if isinstance(value, date):
        return value, None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        d = excel_serial_to_date(value)
        if d:
            return d, None
        return None, f"angka tidak valid sebagai tanggal: {value}"
    if isinstance(value, str):
        s = value.strip()
        if s == "":
            return None, "kosong"
        s_low = s.lower()
        for id_m, en_m in BULAN_ID.items():
            if id_m in s_low:
                s_low = s_low.replace(id_m, en_m)
                break
        for cand in (s, s_low):
            for fmt in EXPLICIT_DATE_FORMATS:
                try:
                    return datetime.strptime(cand.strip(), fmt).date(), None
                except Exception:
                    continue
        try:
            from dateutil import parser as dtparser
            d = dtparser.parse(s, dayfirst=True, fuzzy=False)
            return d.date(), None
        except Exception:
            return None, f"format tanggal tidak dikenali: '{value}'"
    return None, f"tipe data tidak dikenali: {value!r}"


def is_already_valid_date_format(raw):
    if isinstance(raw, (datetime, date)):
        return True
    if isinstance(raw, str) and VALID_DATE_STRING_RE.match(raw.strip()):
        return True
    return False


def try_parse_qty(value):
    if value is None:
        return None, "kosong"
    try:
        if isinstance(value, str):
            v = value.strip().replace(",", "")
            if v == "":
                return None, "kosong"
            f = float(v)
        else:
            f = float(value)
        i = int(round(f))
        if abs(i - f) > 1e-9:
            return None, f"bukan bilangan bulat ({value})"
        if i == 0:
            return None, "qty = 0"
        return i, None
    except Exception:
        return None, f"bukan angka ({value})"


def is_empty(value):
    return value is None or str(value).strip() == ""


# ---------------------------------------------------------------------------
# Load file (xlsx / csv) menjadi DataFrame mentah
# ---------------------------------------------------------------------------
def load_raw_rows(uploaded_file):
    name = uploaded_file.name.lower()
    uploaded_file.seek(0)
    if name.endswith(".csv"):
        df = pd.read_csv(uploaded_file, header=None, dtype=object)
        rows = df.values.tolist()
    else:
        wb = openpyxl.load_workbook(uploaded_file, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return rows


def build_dataframe_from_files(uploaded_files):
    all_records = []
    mapping_report = []
    for f in uploaded_files:
        rows = load_raw_rows(f)
        if not rows:
            continue
        header_idx, score = find_header_row(rows)
        header_vals = rows[header_idx]
        mapping = auto_map_columns(header_vals)
        mapping_report.append({
            "file": f.name, "header_row": header_idx + 1,
            "mapping": mapping, "header_vals": header_vals,
        })
        data_rows = rows[header_idx + 1:]
        for r_i, r in enumerate(data_rows):
            rec = {"_source_file": f.name, "_source_row": header_idx + 2 + r_i}
            for col in COLUMNS:
                idx = mapping.get(col["key"])
                rec[col["key"]] = r[idx] if (idx is not None and idx < len(r)) else None
            if all(rec.get(c["key"]) in (None, "") for c in COLUMNS):
                continue
            all_records.append(rec)
    df = pd.DataFrame(all_records)
    return df, mapping_report


# ---------------------------------------------------------------------------
# Pembersihan otomatis + laporan
# ---------------------------------------------------------------------------
def row_brief(rec):
    return {
        "File Asal": rec.get("_source_file"),
        "Baris Asal": rec.get("_source_row"),
        "Customer Code": rec.get("customer_code"),
        "PO Number": rec.get("po_number"),
        "Store Code": rec.get("store_code"),
        "SKU Code": rec.get("sku_code"),
    }


def clean_and_report(df):
    kept_rows = []
    removed_qty_zero = []
    flagged_missing = []
    flagged_qty_invalid = []
    flagged_date_invalid = []
    date_fixed = []

    for _, rec in df.iterrows():
        rec = rec.to_dict()
        notes = []

        qty_val, qty_err = try_parse_qty(rec.get("qty"))
        if qty_err == "qty = 0":
            item = row_brief(rec)
            item["Qty Asli"] = rec.get("qty")
            removed_qty_zero.append(item)
            continue  # satu-satunya kondisi yang dihapus

        missing_fields = [LABEL_BY_KEY[k] for k in MANDATORY_TEXT_KEYS if is_empty(rec.get(k))]
        if missing_fields:
            note = f"Field kosong: {', '.join(missing_fields)}"
            notes.append(note)
            item = row_brief(rec)
            item["Keterangan"] = note
            flagged_missing.append(item)

        if qty_err:
            note = f"Qty tidak valid ({qty_err}) — nilai asli dipertahankan"
            notes.append(note)
            item = row_brief(rec)
            item["Qty Asli"] = rec.get("qty")
            item["Keterangan"] = note
            flagged_qty_invalid.append(item)

        date_val, date_err = try_parse_date(rec.get("po_date"))
        if date_err:
            note = f"PO Date tidak terbaca ({date_err}) — nilai asli dipertahankan"
            notes.append(note)
            item = row_brief(rec)
            item["PO Date Asli"] = rec.get("po_date")
            item["Keterangan"] = note
            flagged_date_invalid.append(item)
        elif not is_already_valid_date_format(rec.get("po_date")):
            new_str = date_val.strftime("%Y-%m-%d")
            note = f"Format PO Date dirapikan: '{rec.get('po_date')}' → '{new_str}'"
            notes.append(note)
            item = row_brief(rec)
            item["PO Date Asli"] = rec.get("po_date")
            item["PO Date Baru"] = new_str
            date_fixed.append(item)

        rec["_po_date_norm"] = date_val  # None jika tidak terbaca -> nilai asli tetap ditulis
        rec["_qty_norm"] = qty_val       # None jika tidak valid -> nilai asli tetap ditulis
        rec["_catatan"] = "; ".join(notes)
        kept_rows.append(rec)

    clean_df = pd.DataFrame(kept_rows) if kept_rows else pd.DataFrame(
        columns=list(df.columns) + ["_po_date_norm", "_qty_norm", "_catatan"]
    )

    report = {
        "total_awal": len(df),
        "total_akhir": len(clean_df),
        "removed_qty_zero": pd.DataFrame(removed_qty_zero),
        "flagged_missing": pd.DataFrame(flagged_missing),
        "flagged_qty_invalid": pd.DataFrame(flagged_qty_invalid),
        "flagged_date_invalid": pd.DataFrame(flagged_date_invalid),
        "date_fixed": pd.DataFrame(date_fixed),
    }
    return clean_df, report


def build_report_workbook(report):
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = "Ringkasan"
    ws0.append(["Ringkasan Pemrosesan Data", ""])
    ws0.append(["Total baris awal", report["total_awal"]])
    ws0.append(["Dihapus - Qty = 0 (satu-satunya yang dihapus)", len(report["removed_qty_zero"])])
    ws0.append(["Ditandai (tetap ada) - field wajib kosong", len(report["flagged_missing"])])
    ws0.append(["Ditandai (tetap ada) - Qty tidak valid", len(report["flagged_qty_invalid"])])
    ws0.append(["Ditandai (tetap ada) - PO Date tidak terbaca", len(report["flagged_date_invalid"])])
    ws0.append(["Format PO Date dirapikan otomatis", len(report["date_fixed"])])
    ws0.append(["Total baris pada file hasil (final)", report["total_akhir"]])

    sheet_map = {
        "Dihapus_Qty_Nol": report["removed_qty_zero"],
        "Ditandai_Field_Kosong": report["flagged_missing"],
        "Ditandai_Qty_Invalid": report["flagged_qty_invalid"],
        "Ditandai_Tanggal_Invalid": report["flagged_date_invalid"],
        "Format_Tanggal_Dirapikan": report["date_fixed"],
    }
    for sheet_name, sdf in sheet_map.items():
        ws = wb.create_sheet(sheet_name[:31])
        if sdf.empty:
            ws.append(["(tidak ada)"])
            continue
        ws.append(list(sdf.columns))
        for _, r in sdf.iterrows():
            ws.append([("" if pd.isna(v) else v) for v in r.tolist()])

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# Build output files dari template
# ---------------------------------------------------------------------------
def get_template_info(template_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    header_idx, _ = find_header_row(rows)
    header_vals = rows[header_idx]
    mapping = auto_map_columns(header_vals)
    return {
        "sheet_name": ws.title,
        "header_row_1based": header_idx + 1,
        "mapping": mapping,
        "max_row": ws.max_row,
    }


def build_chunk_workbook(template_bytes, template_info, chunk_df):
    wb = openpyxl.load_workbook(io.BytesIO(template_bytes))
    ws = wb[template_info["sheet_name"]]
    header_row = template_info["header_row_1based"]
    mapping = template_info["mapping"]
    has_notes = "_catatan" in chunk_df.columns

    if ws.max_row > header_row:
        ws.delete_rows(header_row + 1, ws.max_row - header_row)

    notes_col_0based = max(mapping.values(), default=-1) + 1  # kolom kosong tepat setelah kolom terakhir template
    if has_notes:
        ws.cell(row=header_row, column=notes_col_0based + 1).value = "Catatan Validasi (otomatis)"
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(notes_col_0based + 1)].width = 45

    for i, (_, row) in enumerate(chunk_df.iterrows()):
        excel_row = header_row + 1 + i
        for col in COLUMNS:
            key = col["key"]
            col_idx_0 = mapping.get(key)
            if col_idx_0 is None:
                continue
            excel_col = col_idx_0 + 1
            cell = ws.cell(row=excel_row, column=excel_col)
            if key == "po_date":
                d = row.get("_po_date_norm")
                if d is not None and not pd.isna(d):
                    cell.value = d
                    cell.number_format = "yyyy-mm-dd"
                else:
                    cell.value = row.get("po_date")
            elif key == "qty":
                q = row.get("_qty_norm")
                cell.value = q if (q is not None and not pd.isna(q)) else row.get("qty")
            else:
                cell.value = row.get(key)
        if has_notes:
            catatan = row.get("_catatan")
            if catatan:
                ws.cell(row=excel_row, column=notes_col_0based + 1).value = catatan

    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.getvalue()


def split_and_build(template_bytes, template_info, df, chunk_size, base_name):
    n = len(df)
    n_chunks = math.ceil(n / chunk_size) if n else 0
    files = {}
    plan = []
    for c in range(n_chunks):
        start = c * chunk_size
        end = min(start + chunk_size, n)
        chunk_df = df.iloc[start:end]
        fname = f"{base_name}_part{c+1}of{n_chunks}_rows{start+1}-{end}.xlsx"
        data_bytes = build_chunk_workbook(template_bytes, template_info, chunk_df)
        files[fname] = data_bytes
        plan.append({"file": fname, "jumlah_baris": end - start})
    return files, plan


def make_zip(files_dict):
    bio = io.BytesIO()
    with zipfile.ZipFile(bio, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in files_dict.items():
            zf.writestr(name, data)
    bio.seek(0)
    return bio.getvalue()


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
st.set_page_config(page_title="DMS Data Splitter", page_icon="📦", layout="wide")
st.title("📦 DMS Data Import Splitter")
st.caption(
    "Upload data PO gabungan → Qty=0 dihapus otomatis, PO Date dirapikan ke yyyy-mm-dd, "
    "baris bermasalah lain tetap dicantumkan dengan catatan → dibagi menjadi beberapa file "
    "xlsx sesuai batas maksimum baris DMS. Tidak ada langkah edit manual."
)

for key in ("result", ):
    if key not in st.session_state:
        st.session_state[key] = None

with st.sidebar:
    st.header("1️⃣ Template DMS")
    tpl_file = st.file_uploader("Upload template resmi (.xlsx)", type=["xlsx"], key="tpl")
    if tpl_file is not None:
        template_bytes = tpl_file.getvalue()
        st.success(f"Menggunakan template: {tpl_file.name}")
    elif os.path.exists(DEFAULT_TEMPLATE_PATH):
        with open(DEFAULT_TEMPLATE_PATH, "rb") as f:
            template_bytes = f.read()
        st.info("Menggunakan template bawaan aplikasi.")
    else:
        template_bytes = None
        st.warning("Belum ada template. Silakan upload template resmi.")

    st.header("2️⃣ Ukuran per file")
    chunk_size = st.number_input("Maksimum baris per file", min_value=1, value=6000, step=100)
    base_name = st.text_input("Nama dasar file output", value="ST_Data_Import")

st.header("3️⃣ Upload Data PO")
data_files = st.file_uploader(
    "Upload satu atau beberapa file data (.xlsx / .csv) yang akan dibagi",
    type=["xlsx", "csv"], accept_multiple_files=True, key="data",
)

run = st.button("🚀 Proses, Bersihkan & Bagi File", type="primary", disabled=not (data_files and template_bytes))

if run:
    df, mapping_report = build_dataframe_from_files(data_files)
    clean_df, report = clean_and_report(df)
    template_info = get_template_info(template_bytes)

    raw_files, raw_plan = split_and_build(template_bytes, template_info, df, int(chunk_size), base_name + "_RAW")
    clean_files, clean_plan = split_and_build(template_bytes, template_info, clean_df, int(chunk_size), base_name + "_CLEAN")
    report_bytes = build_report_workbook(report)

    st.session_state.result = {
        "mapping_report": mapping_report,
        "report": report,
        "raw_files": raw_files, "raw_plan": raw_plan,
        "clean_files": clean_files, "clean_plan": clean_plan,
        "report_bytes": report_bytes,
    }

result = st.session_state.result
if result:
    report = result["report"]

    with st.expander("🔎 Cek pemetaan kolom (otomatis)", expanded=False):
        for rep in result["mapping_report"]:
            st.write(f"**{rep['file']}** (header terdeteksi di baris {rep['header_row']})")
            rows = []
            for col in COLUMNS:
                idx = rep["mapping"].get(col["key"])
                detected = rep["header_vals"][idx] if idx is not None else "❌ tidak ditemukan"
                rows.append({"Kolom Template": col["label"] + (" *" if col["mandatory"] else ""), "Terdeteksi": detected})
            st.dataframe(pd.DataFrame(rows), hide_index=True, use_container_width=True)

    st.header("📋 Ringkasan Hasil Pemrosesan")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total baris awal", report["total_awal"])
    c2.metric("Dihapus (Qty = 0)", len(report["removed_qty_zero"]))
    c3.metric("Ditandai (tetap ada)", len(report["flagged_missing"]) + len(report["flagged_qty_invalid"]) + len(report["flagged_date_invalid"]))
    c4.metric("Baris pada file hasil", report["total_akhir"])

    st.markdown(
        f"- 🗑️ **Qty = 0** — satu-satunya kondisi yang **dihapus**: {len(report['removed_qty_zero'])} baris\n"
        f"- 🏷️ **Field wajib kosong** — tetap dicantumkan + diberi catatan: {len(report['flagged_missing'])} baris\n"
        f"- 🏷️ **Qty tidak valid** (bukan angka/bukan bulat) — tetap dicantumkan + diberi catatan: {len(report['flagged_qty_invalid'])} baris\n"
        f"- 🏷️ **PO Date tidak terbaca** — tetap dicantumkan + diberi catatan: {len(report['flagged_date_invalid'])} baris\n"
        f"- 🛠️ **Format PO Date dirapikan otomatis** ke yyyy-mm-dd: {len(report['date_fixed'])} baris\n\n"
        f"Baris yang ditandai **tidak dihapus** — nilai aslinya tetap ditulis di file, dan kolom "
        f"**'Catatan Validasi'** ditambahkan tepat setelah kolom Qty untuk menjelaskan masalahnya."
    )

    detail_specs = [
        ("🗑️ Baris dihapus — Qty = 0", "removed_qty_zero"),
        ("🏷️ Baris ditandai — Field wajib kosong", "flagged_missing"),
        ("🏷️ Baris ditandai — Qty tidak valid", "flagged_qty_invalid"),
        ("🏷️ Baris ditandai — PO Date tidak terbaca", "flagged_date_invalid"),
        ("🛠️ Baris — format PO Date dirapikan (asli → baru)", "date_fixed"),
    ]
    for title, key in detail_specs:
        sdf = report[key]
        with st.expander(f"{title} ({len(sdf)} baris)", expanded=False):
            if sdf.empty:
                st.write("Tidak ada.")
            else:
                st.dataframe(sdf, hide_index=True, use_container_width=True)

    st.download_button(
        "⬇️ Download Laporan Lengkap (Excel)", data=result["report_bytes"],
        file_name=f"{base_name}_Laporan_Pembersihan.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.divider()
    st.header("4️⃣ Download Hasil")

    colA, colB = st.columns(2)
    with colA:
        st.subheader("Versi 1: Split Saja (data asli, tanpa dibersihkan)")
        st.caption("Data original langsung dibagi ke template, tidak ada baris yang dihapus/diubah.")
        st.dataframe(pd.DataFrame(result["raw_plan"]), hide_index=True, use_container_width=True)
        st.download_button(
            "⬇️ Download Versi Split Saja (ZIP)",
            data=make_zip(result["raw_files"]),
            file_name=f"{base_name}_split_saja.zip", mime="application/zip",
        )
        with st.expander("Download satuan (versi split saja)"):
            for fname, data in result["raw_files"].items():
                st.download_button(f"⬇️ {fname}", data=data, file_name=fname, key=f"raw_{fname}")

    with colB:
        st.subheader("Versi 2: Sudah Dibersihkan + Split")
        st.caption("Qty=0 sudah dihapus. Baris lain yang bermasalah tetap ada, plus kolom Catatan Validasi & PO Date dirapikan ke yyyy-mm-dd bila terbaca.")
        st.dataframe(pd.DataFrame(result["clean_plan"]), hide_index=True, use_container_width=True)
        st.download_button(
            "⬇️ Download Versi Sudah Dibersihkan (ZIP)",
            data=make_zip(result["clean_files"]),
            file_name=f"{base_name}_clean_split.zip", mime="application/zip", type="primary",
        )
        with st.expander("Download satuan (versi sudah dibersihkan)"):
            for fname, data in result["clean_files"].items():
                st.download_button(f"⬇️ {fname}", data=data, file_name=fname, key=f"clean_{fname}")
else:
    st.info("Upload template & data, lalu klik tombol proses di atas.")
